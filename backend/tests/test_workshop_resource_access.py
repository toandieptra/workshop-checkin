"""HTTP authorization tests with real ORM queries and authorization dependencies.

SQLite supplies isolated relational fixtures; the small awaitable adapter lets
these request tests exercise the same query paths without a running Postgres.
"""
import io
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.ext.compiler import compiles
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.auth.dependencies import current_user
from app.db import Base, get_db
from app.models import (AdminUser, Guest, GuestNote, RegistrationForm,
                        RegistrationFormWorkshop, RolePermission, Workshop,
                        WorkshopMember, ZaloDelivery, ZaloDeliveryItem, ZaloTemplate)
from app.routers import admin_workshops, auth, checkin, guests, import_export, registration_forms, search, workshops, zalo_messages, zbs


@compiles(JSONB, "sqlite")
def _compile_jsonb(_element, _compiler, **_kwargs):
    return "JSON"


class AwaitableSession:
    def __init__(self, session):
        self.session = session

    async def execute(self, statement, *args, **kwargs):
        return self.session.execute(statement, *args, **kwargs)

    async def scalar(self, statement, *args, **kwargs):
        return self.session.scalar(statement, *args, **kwargs)

    async def get(self, model, key):
        return self.session.get(model, key)

    def add(self, value):
        self.session.add(value)

    async def flush(self):
        self.session.flush()

    async def commit(self):
        self.session.commit()

    async def refresh(self, value):
        self.session.refresh(value)

    async def rollback(self):
        self.session.rollback()

    async def delete(self, value):
        self.session.delete(value)


@pytest.fixture
def ctx():
    engine = create_engine("sqlite://", poolclass=StaticPool, connect_args={"check_same_thread": False})

    @event.listens_for(engine, "connect")
    def install_uuid(connection, _record):
        connection.create_function("uuid_generate_v4", 0, lambda: uuid.uuid4().hex)

    Base.metadata.create_all(engine)
    session = Session(engine, expire_on_commit=False)
    db = AwaitableSession(session)
    user = AdminUser(id=uuid.uuid4(), email="scoped@example.com", name="Scoped user", role="user", is_active=True, permission_overrides={})
    ws = [Workshop(id=uuid.uuid4(), name=f"Workshop {name}", slug=f"workshop-{name.lower()}", status="published") for name in "ABC"]
    records = [Guest(id=uuid.uuid4(), workshop_id=workshop.id, full_name=f"Guest {name}", registration_status="pending") for workshop, name in zip(ws, "ABC")]
    members = [WorkshopMember(user_id=user.id, workshop_id=ws[0].id, role="editor"), WorkshopMember(user_id=user.id, workshop_id=ws[1].id, role="viewer")]
    session.add_all([user, *ws, *records, *members])
    session.commit()
    app = FastAPI()
    for module in (admin_workshops, auth, workshops, guests, checkin, search, import_export, registration_forms, zbs, zalo_messages):
        app.include_router(module.router)

    async def database():
        try:
            yield db
        finally:
            session.rollback()

    async def account():
        return user

    app.dependency_overrides[get_db] = database
    app.dependency_overrides[current_user] = account
    with TestClient(app) as client:
        yield SimpleNamespace(app=app, client=client, db=db, session=session, user=user, workshops=ws, guests=records, members=members)
    session.close()
    engine.dispose()


def _grant(ctx, index, role="admin"):
    ctx.members[index].role = role
    ctx.session.commit()


def _form(ctx, indices):
    form = RegistrationForm(id=uuid.uuid4(), workshop_id=ctx.workshops[indices[0]].id, token=uuid.uuid4().hex)
    ctx.session.add(form)
    ctx.session.add_all([RegistrationFormWorkshop(form_id=form.id, workshop_id=ctx.workshops[i].id) for i in indices])
    ctx.session.commit()
    return form


def _delivery(ctx, indices):
    delivery = ZaloDelivery(id=uuid.uuid4(), template_name="History", content_blocks=[{"type": "text", "text": "Hello"}], idempotency_key=uuid.uuid4().hex, recipient_count=len(indices))
    ctx.session.add(delivery)
    ctx.session.add_all([ZaloDeliveryItem(id=uuid.uuid4(), delivery_id=delivery.id, guest_id=ctx.guests[i].id if i is not None else None,
        block_position=0, block_payload={"type": "text", "text": "Hello"}, status="failed", recipient_id="recipient") for i in indices])
    ctx.session.commit()
    return delivery


def test_metadata_search_and_direct_guest_access_are_scoped(ctx):
    client = ctx.client
    assert {row["id"] for row in client.get("/api/workshops").json()} == {str(ctx.workshops[i].id) for i in (0, 1)}
    assert {row["id"] for row in client.get("/api/search/guests", params={"q": "Guest"}).json()} == {str(ctx.guests[i].id) for i in (0, 1)}
    assert client.get(f"/api/guests/{ctx.guests[1].id}").status_code == 200
    assert client.patch(f"/api/guests/{ctx.guests[0].id}", json={"full_name": "Allowed edit"}).status_code == 200
    assert client.patch(f"/api/guests/{ctx.guests[1].id}", json={"full_name": "Forbidden edit"}).status_code == 403
    assert client.get(f"/api/guests/{ctx.guests[2].id}").status_code == 403
    assert client.get(f"/api/workshops/{ctx.workshops[2].id}").status_code == 403
    assert client.get("/api/search/guests", params={"q": "Guest", "workshop_id": str(ctx.workshops[2].id)}).status_code == 403
    assert ctx.guests[1].full_name == "Guest B"


def test_scoped_workshop_admin_cannot_use_global_creation(ctx):
    _grant(ctx, 0)
    assert ctx.client.post("/api/workshops", json={"name": "Denied"}).status_code == 403
    assert ctx.client.patch(f"/api/workshops/{ctx.workshops[0].id}", json={"name": "Allowed"}).status_code == 200
    assert ctx.client.patch(f"/api/workshops/{ctx.workshops[1].id}", json={"name": "Denied"}).status_code == 403


def test_new_workshop_grants_only_its_creator(ctx):
    ctx.user.permission_overrides = {"workshops.write": True}
    ctx.session.commit()
    response = ctx.client.post("/api/workshops", json={"name": "New workshop", "slug": "new-workshop"})
    assert response.status_code == 201, response.text
    created_id = uuid.UUID(response.json()["id"])
    member = ctx.session.get(WorkshopMember, (ctx.user.id, created_id))
    assert member is not None and member.role == "user"


@pytest.mark.parametrize("role", ["admin", "super_admin"])
def test_global_administrators_access_new_workshops_with_functional_overrides(ctx, role):
    ctx.user.role = role
    ctx.session.commit()
    assert len(ctx.client.get("/api/workshops").json()) == 3
    ctx.user.permission_overrides = {"guests.write": False}
    ctx.session.commit()
    assert ctx.client.patch(f"/api/guests/{ctx.guests[2].id}", json={"full_name": "Denied"}).status_code == 403


def test_membership_revocation_and_overrides_take_effect_next_request(ctx):
    assert ctx.client.get(f"/api/guests/{ctx.guests[0].id}").status_code == 200
    ctx.session.delete(ctx.members[0])
    ctx.user.permission_overrides = {"guests.read": True, "guests.write": True}
    ctx.session.commit()
    assert ctx.client.get(f"/api/guests/{ctx.guests[0].id}").status_code == 403
    assert ctx.client.patch(f"/api/guests/{ctx.guests[1].id}", json={"full_name": "Override"}).status_code == 200
    ctx.user.is_active = False
    ctx.session.commit()
    assert ctx.client.get(f"/api/guests/{ctx.guests[1].id}").status_code == 403


def test_resource_ids_resolve_workshop_before_notes_checkin_import_or_zbs(ctx):
    client = ctx.client
    guest = ctx.guests[2]
    workshop = ctx.workshops[2]
    cases = [
        ("get", f"/api/guests/{guest.id}/notes", None),
        ("post", f"/api/guests/{guest.id}/notes", {"content": "Denied"}),
        ("post", f"/api/guests/{guest.id}/checkin", {}),
        ("post", f"/api/guests/{guest.id}/uncheckin", None),
        ("get", f"/api/guests/{guest.id}/qr-info", None),
        ("post", f"/api/checkin/reset?guest_id={guest.id}", None),
        ("get", f"/api/checkin/logs?workshop_id={workshop.id}", None),
        ("get", f"/api/workshops/{workshop.id}/zbs-status", None),
        ("post", f"/api/zbs/guests/{guest.id}/send/registration_confirmation", None),
        ("get", f"/api/workshops/{workshop.id}/landing-page/content", None),
    ]
    for method, url, payload in cases:
        assert client.request(method, url, json=payload).status_code == 403, url
    response = client.post(f"/api/workshops/{workshop.id}/import", files={"file": ("guests.csv", b"full_name\nDenied")})
    assert response.status_code == 403
    assert guest.checkin_status == "not_checked_in"
    assert ctx.session.scalar(select(GuestNote)) is None


def test_shared_forms_require_every_old_and_new_workshop(ctx):
    private = _form(ctx, [0])
    shared = _form(ctx, [0, 2])
    assert [row["id"] for row in ctx.client.get("/api/registration-forms").json()] == [str(private.id)]
    linked = ctx.client.get(f"/api/workshops/{ctx.workshops[0].id}/registration-forms")
    assert [row["id"] for row in linked.json()] == [str(private.id)]
    embedded = ctx.client.get(f"/api/workshops/{ctx.workshops[0].id}").json()["registration_forms"]
    assert [row["id"] for row in embedded] == [str(private.id)]
    assert ctx.client.patch(f"/api/registration-forms/{shared.id}", json={"greeting": "Denied", "workshop_ids": [str(ctx.workshops[0].id)]}).status_code == 403
    assert ctx.client.patch(f"/api/registration-forms/{private.id}", json={"greeting": "Denied", "workshop_ids": [str(ctx.workshops[0].id), str(ctx.workshops[2].id)]}).status_code == 403
    assert ctx.client.delete(f"/api/registration-forms/{shared.id}").status_code == 403
    assert shared.greeting is None and private.greeting is None


def test_legacy_form_falls_back_to_primary_workshop(ctx):
    form = RegistrationForm(id=uuid.uuid4(), workshop_id=ctx.workshops[2].id, token=uuid.uuid4().hex)
    ctx.session.add(form)
    ctx.session.commit()
    assert ctx.client.get("/api/registration-forms").json() == []
    assert ctx.client.patch(f"/api/registration-forms/{form.id}", json={"greeting": "Denied"}).status_code == 403


def _export_names(response):
    assert response.status_code == 200, response.text
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    rows = list(workbook.active.values)
    name_column = list(rows[0]).index("full_name")
    return {row[name_column] for row in rows[1:]}


def test_exports_scope_all_and_reject_explicit_ungranted_workshops(ctx):
    ctx.user.permission_overrides = {"guests.export": True}
    ctx.session.commit()
    assert _export_names(ctx.client.get("/api/export/guests")) == {"Guest A", "Guest B"}
    assert ctx.client.get("/api/export/guests", params={"workshop_ids": f"{ctx.workshops[0].id},{ctx.workshops[2].id}"}).status_code == 403
    assert _export_names(ctx.client.get("/api/export/guests", params={"export_context": "reports"})) == {"Guest A"}
    assert ctx.client.get("/api/export/guests", params={"export_context": "reports", "workshop_id": str(ctx.workshops[1].id)}).status_code == 403


def test_custom_role_cannot_mix_report_read_and_export_between_workshops(ctx):
    ctx.session.add_all([RolePermission(role="report_reader", label="Reader", permissions=["reports.read"]), RolePermission(role="report_exporter", label="Exporter", permissions=["reports.export"])])
    ctx.members[0].role, ctx.members[1].role = "report_reader", "report_exporter"
    ctx.session.commit()
    assert _export_names(ctx.client.get("/api/export/guests", params={"export_context": "reports"})) == set()
    assert ctx.client.get("/api/export/guests", params={"export_context": "reports", "workshop_id": str(ctx.workshops[1].id)}).status_code == 403


def test_zalo_mixed_recipients_reject_before_lookup_or_creation(ctx, monkeypatch):
    _grant(ctx, 0)
    template = ZaloTemplate(id=uuid.uuid4(), name="Template", content_blocks=[{"type": "text", "text": "Hello"}], status="active")
    ctx.session.add(template)
    ctx.session.commit()
    resolver, creator = AsyncMock(), AsyncMock()
    monkeypatch.setattr(zalo_messages, "resolve_recipients", resolver)
    monkeypatch.setattr(zalo_messages, "create_delivery", creator)
    payload = {"template_id": str(template.id), "guest_ids": [str(ctx.guests[0].id), str(ctx.guests[2].id)]}
    assert ctx.client.post("/api/zalo/preflight", json=payload).status_code == 403
    assert ctx.client.post("/api/zalo/batches", json={**payload, "idempotency_key": str(uuid.uuid4())}).status_code == 403
    resolver.assert_not_awaited()
    creator.assert_not_awaited()
    assert ctx.session.scalar(select(ZaloDelivery)) is None


def test_zalo_history_filters_mixed_unresolved_and_reauthorizes_retries(ctx, monkeypatch):
    _grant(ctx, 0)
    own, mixed, unresolved = _delivery(ctx, [0]), _delivery(ctx, [0, 2]), _delivery(ctx, [None])
    response = ctx.client.get("/api/zalo/deliveries")
    assert response.status_code == 200, response.text
    assert [row["id"] for row in response.json()] == [str(own.id)]
    assert response.json()[0]["items"][0]["workshop_id"] == str(ctx.workshops[0].id)
    for delivery in (mixed, unresolved):
        assert ctx.client.get(f"/api/zalo/deliveries/{delivery.id}").status_code == 403
    retry_quota, resolve = AsyncMock(), AsyncMock()
    monkeypatch.setattr(zalo_messages, "reopen_quota", retry_quota)
    monkeypatch.setattr(zalo_messages, "resolve_recipient", resolve)
    own_item_in_mixed = next(item for item in mixed.items if item.guest_id == ctx.guests[0].id)
    assert ctx.client.post(f"/api/zalo/delivery-items/{own_item_in_mixed.id}/retry").status_code == 403
    assert ctx.client.post(f"/api/zalo/delivery-items/{own_item_in_mixed.id}/refresh").status_code == 403
    retry_quota.assert_not_awaited()
    resolve.assert_not_awaited()


def test_zalo_idempotency_hit_reauthorizes_whole_existing_delivery(ctx, monkeypatch):
    _grant(ctx, 0)
    existing = _delivery(ctx, [0, 2])
    existing.idempotency_key = str(uuid.uuid4())
    template = ZaloTemplate(id=uuid.uuid4(), name="Template", content_blocks=[{"type": "text", "text": "Hello"}], status="active")
    ctx.session.add(template)
    ctx.session.commit()
    creator = AsyncMock()
    monkeypatch.setattr(zalo_messages, "create_delivery", creator)
    response = ctx.client.post("/api/zalo/send", json={"template_id": str(template.id), "guest_id": str(ctx.guests[0].id), "idempotency_key": existing.idempotency_key})
    assert response.status_code == 403
    creator.assert_not_awaited()


def test_workshop_admin_endpoints_require_authentication(ctx):
    del ctx.app.dependency_overrides[current_user]
    assert ctx.client.get("/api/workshops").status_code == 401
    assert ctx.client.get(f"/api/workshops/{ctx.workshops[0].id}").status_code == 401
    assert ctx.client.get("/api/workshops/meta/branches").status_code == 401


def test_module_workshop_selectors_use_requested_functional_permission(ctx):
    ctx.session.add(RolePermission(role="report_only", label="Report", permissions=["reports.read"]))
    ctx.members[0].role = "report_only"
    ctx.session.commit()
    assert [row["id"] for row in ctx.client.get("/api/workshops").json()] == [str(ctx.workshops[1].id)]
    rows = ctx.client.get("/api/workshops", params={"permission": "reports.read"}).json()
    assert {row["id"] for row in rows} == {str(ctx.workshops[0].id), str(ctx.workshops[1].id)}
    assert all(row["registration_forms"] == [] for row in rows)
    assert ctx.client.get(f"/api/workshops/{ctx.workshops[0].id}").status_code == 403
    assert ctx.client.get("/api/workshops", params={"permission": "users.manage"}).status_code == 422


def test_report_only_role_can_read_report_data_but_not_use_guest_module(ctx):
    ctx.session.add(RolePermission(role="report_only", label="Report", permissions=["reports.read", "reports.export"]))
    ctx.members[0].role = "report_only"
    ctx.session.commit()
    path = f"/api/workshops/{ctx.workshops[0].id}/guests"
    assert ctx.client.get(path).status_code == 403
    response = ctx.client.get(path, params={"permission": "reports.read"})
    assert response.status_code == 200, response.text
    assert [row["id"] for row in response.json()] == [str(ctx.guests[0].id)]
    assert ctx.client.get(path, params={"permission": "users.manage"}).status_code == 422
    assert ctx.client.get(f"/api/workshops/{ctx.workshops[2].id}/guests", params={"permission": "reports.read"}).status_code == 403
    assert _export_names(ctx.client.get("/api/export/guests", params={"export_context": "reports"})) == {"Guest A"}


def test_guest_only_role_can_enumerate_its_guest_workshops(ctx):
    ctx.session.add(RolePermission(role="guest_only", label="Guest", permissions=["guests.read"]))
    ctx.members[0].role = "guest_only"
    ctx.session.commit()
    assert [row["id"] for row in ctx.client.get("/api/workshops").json()] == [str(ctx.workshops[1].id)]
    response = ctx.client.get("/api/workshops", params={"permission": "guests.read"})
    assert {row["id"] for row in response.json()} == {str(ctx.workshops[0].id), str(ctx.workshops[1].id)}


def test_auth_payload_keeps_global_permissions_separate_from_workshop_roles(ctx):
    _grant(ctx, 0, "admin")
    response = ctx.client.get("/api/auth/me")
    assert response.status_code == 200, response.text
    assert response.json()["permissions"] == []
    access = {item["workshop_id"]: item for item in response.json()["workshop_access"]}
    assert "workshops.delete" in access[str(ctx.workshops[0].id)]["permissions"]
    assert "workshops.delete" not in access[str(ctx.workshops[1].id)]["permissions"]
    assert all("users.manage" not in item["permissions"] for item in access.values())


def test_membership_replacement_persists_and_rejects_invalid_batch_atomically(ctx):
    ctx.user.role = "super_admin"
    target = AdminUser(id=uuid.uuid4(), email="inactive@example.test", role="user", is_active=False, permission_overrides={})
    ctx.session.add_all([target, RolePermission(role="viewer", label="Viewer", permissions=["guests.read"])])
    ctx.session.commit()
    path = f"/api/admin/workshops/{ctx.workshops[0].id}/members"
    assignment = {"user_id": str(target.id), "role": "viewer"}
    response = ctx.client.put(path, json={"members": [assignment]})
    assert response.status_code == 200, response.text
    assert response.json()["members"] == [assignment]
    assert response.json()["inherited_user_ids"] == [str(ctx.user.id)]
    assert target.is_active is False
    assert ctx.client.put(path, json={"members": [{**assignment, "role": "unknown"}]}).status_code == 400
    assert ctx.client.get(path).json()["members"] == [assignment]
    assert ctx.client.put(path, json={"members": []}).status_code == 200
    assert ctx.session.get(WorkshopMember, (target.id, ctx.workshops[0].id)) is None


def test_zalo_guest_status_query_filters_entire_mixed_delivery(ctx):
    _grant(ctx, 0)
    template = ZaloTemplate(id=uuid.uuid4(), name="Template", content_blocks=[{"type": "text", "text": "Hello"}], status="active")
    own, mixed = _delivery(ctx, [0]), _delivery(ctx, [0, 2])
    ctx.session.add(template)
    own.template_id = mixed.template_id = template.id
    ctx.session.commit()
    response = ctx.client.get("/api/zalo/guest-send-statuses", params={
        "workshop_id": str(ctx.workshops[0].id), "template_id": str(template.id),
    })
    assert response.status_code == 200, response.text
    assert len(response.json()) == 1
    assert response.json()[0]["delivery_id"] == str(own.id)
